from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path("output/docx/kong-weipeng-job-fair-interview-prep-cn.docx")
FONT = "Arial Unicode MS"
INK = "000000"
MUTED = "4B5563"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
GRID = "D9D9D9"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), FONT)


def set_style_font(style, size, bold=False, color=INK, space_after=6, line_spacing=1.35):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_after = Pt(space_after)
    style.paragraph_format.line_spacing = line_spacing
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def border_cell(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def no_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)


def table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, 8, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(text))
    return p


def add_question(doc, question, answer, boundary=None):
    p = doc.add_paragraph(style="Question")
    set_run_font(p.add_run(question), bold=True, color=INK)
    add_body(doc, "回答重点：" + answer, "回答重点：")
    if boundary:
        add_body(doc, "不要说：" + boundary, "不要说：")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    head = table.rows[0]
    table_header(head)
    for index, text in enumerate(headers):
        cell = head.cells[index]
        cell.width = widths[index]
        shade(cell, NAVY)
        border_cell(cell)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, 9.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        no_row_split(row)
        for col_index, text in enumerate(values):
            cell = row.cells[col_index]
            cell.width = widths[col_index]
            shade(cell, PALE_BLUE if row_index % 2 == 0 else "FFFFFF")
            border_cell(cell)
            cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(text)
            set_run_font(run, 9.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    set_style_font(styles["Normal"], 10.5, color=INK, space_after=5)
    set_style_font(styles["Body Text"], 10.5, color=INK, space_after=5)
    set_style_font(styles["Title"], 22, bold=True, color=INK, space_after=6, line_spacing=1.1)
    set_style_font(styles["Subtitle"], 11, color=MUTED, space_after=14, line_spacing=1.25)
    set_style_font(styles["Heading 1"], 15, bold=True, color=INK, space_after=7, line_spacing=1.15)
    set_style_font(styles["Heading 2"], 12, bold=True, color=INK, space_after=5, line_spacing=1.2)
    set_style_font(styles["List Bullet"], 10.3, color=INK, space_after=3, line_spacing=1.25)
    question = styles.add_style("Question", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(question, 10.8, bold=True, color=INK, space_after=2, line_spacing=1.25)
    question.paragraph_format.space_before = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("孔维鹏  招聘会面试准备  第 ")
    set_run_font(run, 8, color=MUTED)
    add_page_number(footer)
    run = footer.add_run(" 页")
    set_run_font(run, 8, color=MUTED)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("国内产品体验与智能硬件求职知识准备")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("用于深圳海归招聘会和后续产品体验岗位面试  2026年9月")
    add_body(
        doc,
        "这份材料服务于三类岗位：AI和科技产品、智能硬件产品体验、智能汽车HMI。它不替代作品集，而是帮助我把项目经历换成国内招聘中常用且可核实的表达，并提前准备面试中会追问的方法、边界和细节。",
    )
    add_body(
        doc,
        "使用时先看术语对照，再按目标企业复习对应方向。回答项目问题时始终按“场景、发现、方案、验证、当前状态”组织，不用没有做过的数据或结果补强故事。",
    )

    add_heading(doc, "国内招聘中的专业表达", 1)
    add_body(doc, "下面的说法可以直接用于简历、现场介绍和面试。左侧是研究或作品集中的原始表述，右侧是更容易被国内产品、设计和研发团队理解的表达。")
    add_table(
        doc,
        ["原始说法", "国内常用表达", "使用边界"],
        [
            ["感性工学", "用户主观体验与行为数据研究", "说明自己会结合主观评价、眼动或心率等信息分析体验差异。不要把它说成心理诊断。"],
            ["人间情报设计", "人机交互和用户体验研究方向", "首次出现可保留日文专业名，后面统一说人机交互或交互设计。"],
            ["感性评价", "主观体验评价", "说清评价对象，例如提示是否易懂、是否有压迫感、是否愿意使用。"],
            ["概念模拟", "概念原型或场景原型", "视觉示意、视频、Figma流程都属于概念原型，不等于真实产品上线。"],
            ["验证", "概念调研、可用性测试或原型验证", "先说具体方法。问卷不能直接证明真实长期使用效果。"],
            ["AI共创", "人机协同创作或AI辅助创作", "重点是用户如何参与关键判断，不要泛称“AI更懂用户”。"],
            ["判断分担", "用户参与关键判断和决策权设置", "解释哪些内容由AI建议，哪些内容必须由用户确认或补充。"],
            ["眼动追踪", "眼动测试或注意力研究", "明确是已完成、正在计划，还是作为后续验证方案。"],
            ["用AI做小程序", "借助AI开发工具完成可用原型并上线", "准备说明自己做了什么：需求、页面、流程、测试、发布。不要暗示独立完成复杂后端。"],
        ],
        [Cm(3.3), Cm(5.0), Cm(8.0)],
    )

    add_heading(doc, "三类岗位的能力重点", 1)
    add_table(
        doc,
        ["方向", "国内岗位常见名称", "招聘方想确认什么", "我最该拿出的项目"],
        [
            ["科技产品", "产品经理 体验产品经理 AI产品经理 用户体验设计", "能否从模糊场景中提出需求，定义功能和交互，并把原型做成可讨论的方案。", "AI共创界面、好久没吃小程序、FrameTrace"],
            ["智能硬件", "智能硬件产品经理 产品设计师 交互设计师 用户体验研究员", "是否理解实体设备、界面、灯光或传感器如何配合，以及概念如何落到可体验的原型。", "冰箱物品日期与位置管理、AI共创界面、FrameTrace"],
            ["智能汽车HMI", "智能座舱产品经理 HMI设计师 体验产品经理 交互设计师", "是否能在驾驶注意、安全约束和用户感受之间处理信息层级，知道什么不能打扰驾驶员。", "智能驾驶HMI、AI共创界面、FrameTrace"],
        ],
        [Cm(2.3), Cm(4.3), Cm(6.7), Cm(3.0)],
    )
    add_body(doc, "现场递简历前先问岗位属于哪一类。荣耀、视源这类软硬件业务并存的企业，需要先确认部门；海尔优先讲冰箱，汽车企业优先讲HMI，互联网和AI产品团队优先讲AI共创与小程序。")

    add_heading(doc, "产品岗位需要掌握的基础语言", 1)
    add_heading(doc, "从场景到方案", 2)
    add_bullet(doc, "用户场景：谁在什么时间、什么环境下，要完成什么事。不要只描述功能，例如“查找食材”要补充用户为何会忘记位置、翻找的成本是什么。")
    add_bullet(doc, "用户问题：场景中的具体阻碍。产品问题需要可观察，例如信息不在眼前、记录成本高、用户不知道系统正在关注什么。")
    add_bullet(doc, "需求优先级：先判断是否高频、影响是否大、是否能被当前方案解决。面试中可以说“我会先区分高频痛点和展示性功能”。")
    add_bullet(doc, "功能定义：功能入口、用户输入、系统反馈、下一步动作和异常情况。面试官常用“流程是否闭环”追问这里。")
    add_bullet(doc, "交互原型：用于讨论流程和反馈方式的材料。Figma页面、视频模拟、Blender场景图、实体模型都可以成为不同层级的原型。")
    add_bullet(doc, "PRD：产品需求文档。即使没有写过完整PRD，也要能口头讲清目标用户、使用场景、核心流程、规则、边界和验收方式。")

    add_heading(doc, "用户研究和体验验证", 2)
    add_bullet(doc, "探索性调研：用于理解用户行为、需求和语言。可用访谈、开放题问卷、竞品观察。它适合回答“问题是否存在”。")
    add_bullet(doc, "概念调研：把方案通过图、视频或原型展示给受访者，了解第一印象、理解难点和使用意愿。它适合筛选方向，不能替代长期真实使用数据。")
    add_bullet(doc, "可用性测试：让用户完成具体任务，观察是否能理解入口、完成流程和处理错误。常看任务完成情况、错误点、完成时间和主观感受。")
    add_bullet(doc, "主观体验评价：让用户表达清晰度、压力感、愉悦感、信任感或掌控感。题目必须与项目目标对应，不能只问“喜不喜欢”。")
    add_bullet(doc, "眼动测试：用于观察注意区域、停留时长、视线切换和漏看信息。它只说明注意分配，不直接等于喜欢或理解。")
    add_bullet(doc, "样本与数据：有明确样本数就如实说明；没有最终有效样本数时，说“目前完成前期概念调研，正在整理有效样本”，不要补造数字。")

    add_heading(doc, "AI产品面试知识", 1)
    add_body(doc, "AI产品面试重点不是背模型名，而是解释AI在某个场景中应该帮用户做什么、不能替用户做什么，以及如何让用户发现和纠正AI的误解。")
    add_table(
        doc,
        ["概念", "面试中可用的理解", "与AI共创项目的对应"],
        [
            ["大语言模型", "根据上下文预测和生成文本，擅长归纳、改写、对话与方案发散，但可能误解模糊词或生成不可靠内容。", "“高级感”这类词容易被按常见语义解释，忽略个人标准。"],
            ["提示词", "用户给模型的任务说明和上下文。产品设计不应把所有责任交给用户写提示词，而要通过界面帮助用户补充关键信息。", "在AI解释旁设置确认和调整入口，让用户决定是否补充自己的理解。"],
            ["人机协同", "AI负责发散、整理、提出候选方案，用户保留价值判断、取舍和最终确认。", "项目不是让AI替人创作，而是让用户看见并参与容易被忽略的判断。"],
            ["可控性", "用户知道系统依据什么生成，能修改输入、查看关键假设、撤销或继续迭代。", "把模糊词的解释放到可见位置，降低黑箱感。"],
            ["幻觉与误解", "模型可能生成听起来合理但不符合事实或用户意图的内容。产品需要提示风险，并给用户校正入口。", "项目关注的是“意图误解”，不是技术错误率。"],
            ["AI原型工作流", "用AI工具辅助页面、逻辑、文案或小程序开发，并由设计者检查结果、补足业务规则和异常情况。", "好久没吃项目可说明从概念调研到可用小程序的落地过程。"],
        ],
        [Cm(3.0), Cm(7.5), Cm(5.8)],
    )
    add_question(doc, "你认为AI产品最重要的体验问题是什么", "先看用户是否知道系统在做什么、依据什么做，再看用户是否能修正结果。不同场景权重不同：创作场景重视表达和可控性，生活服务场景重视效率和可靠性。", "不要回答“让AI更像人”或“提示词写好就可以”。")
    add_question(doc, "你会怎么评价AI共创界面是否有效", "先比较用户是否能理解AI对模糊词的解释，再观察用户是否愿意补充个人标准。后续可结合眼动比较不同界面中用户是否注意到解释区和调整入口。", "不要把计划中的眼动结果说成已经得出的结论。")

    add_heading(doc, "智能硬件面试知识", 1)
    add_body(doc, "智能硬件岗位会问得比UI更具体。需要把“页面怎么画”扩展到设备如何被发现、状态如何被看见、出错时用户如何处理，以及硬件限制是否被考虑。")
    add_bullet(doc, "输入与输出：输入可以是摄像头、触控、扫码或手机操作；输出可以是屏幕、灯光、声音、震动或手机消息。说明每一种反馈负责解决什么问题。")
    add_bullet(doc, "状态可见性：用户需要知道系统是否识别成功、物品记录在哪里、提醒为什么出现。冰箱项目的外部灯条和内部区域灯光属于状态和位置反馈。")
    add_bullet(doc, "跨设备协同：冰箱本体适合快速定位和取用，手机适合细节查看与长期管理。面试中要说清每个端的任务，不要把所有功能都堆到一个屏幕里。")
    add_bullet(doc, "实体原型：1/4冰箱模型证明的是造型、交互位置和灯光指引能否被体验，不代表制冷、识别精度或量产可靠性已验证。")
    add_bullet(doc, "ID和交互的关系：工业设计关注产品形态、结构和使用方式；交互设计关注用户如何输入、理解和反馈。你的优势是能把两者放在同一使用场景里讨论。")
    add_bullet(doc, "量产意识：概念项目可以谈未来需要验证的安装位置、清洁维护、成本、隐私、误识别和断网情况。不要假装已经解决。")
    add_question(doc, "为什么冰箱需要中央摄像和触控终端", "它的作用是把记录与查询放到使用现场，并通过可磁吸移动的形式贴近不同操作位置；外部和内部灯光负责低成本的快速指引，手机负责更细的记录和管理。", "不要说摄像头可以自动识别所有食材，除非你真的做过识别模型和准确率测试。")
    add_question(doc, "这个冰箱项目离真实产品还有多远", "已经完成场景模拟和带电控的1/4模型，用于验证造型、灯光指引与基本操作。要落地还需要继续验证识别准确性、食材录入成本、隐私处理、硬件耐用性和量产成本。", "不要把概念模型说成产品样机或试生产。")

    add_heading(doc, "智能汽车HMI面试知识", 1)
    add_body(doc, "汽车HMI面试会特别看你有没有安全边界。你的项目优势不是增加一个HUD，而是讨论自动驾驶时如何让驾驶员知道系统关注了什么，并维持对道路和驾驶的参与感。")
    add_bullet(doc, "驾驶任务优先：信息展示不能抢占驾驶员的注意。越紧急的信息越应该短、明确、位置稳定，且不能在同一时刻堆叠多个强提示。")
    add_bullet(doc, "风险分级：需要区分“系统正在关注”“需要用户留意”“需要用户尽快接管”。不同等级对应不同的文案、位置、持续时间和提示强度。")
    add_bullet(doc, "注意维持：自动化容易让人降低观察路况的主动性。设计目标可以是让人理解风险来源，而不是持续制造紧张。")
    add_bullet(doc, "可解释性：当系统提示某个车辆或行人，驾驶员应能理解提示对象和原因。解释不等于展示全部传感器信息，而是展示与当前决策相关的内容。")
    add_bullet(doc, "驾驶价值：高价值车型用户可能希望保留操控感。你的项目可以讲成“在安全前提下，用适度、少量、有对象的提示让驾驶保持参与感”。")
    add_bullet(doc, "自动驾驶分级：面试中至少要能区分辅助驾驶和条件自动驾驶。L2阶段驾驶员仍需要持续承担驾驶责任；无论谈什么界面，都要先承认法规、产品功能和道路条件的限制。")
    add_question(doc, "为什么不用普通的红色感叹号提示", "感叹号只告诉用户有风险，不能告诉用户风险在哪里、为什么重要。我的方案会把提示绑定到具体对象，例如可能突然进入道路的行人或可能并线的车辆，同时按风险强度控制数量和表达方式。", "不要说娱乐化提示一定能提高安全性。应说它是用于提升注意和理解的概念，需要进一步验证。")
    add_question(doc, "娱乐化文案会不会影响驾驶安全", "会有风险，所以它不能出现在高风险、需要立即接管的时刻。更适合低到中等风险的注意维持场景，并且必须短、单一、与对象绑定。高风险场景仍要使用清晰、标准化的提示。", "不要把赛道游戏化或情绪化提示直接套到公共道路。")

    add_heading(doc, "项目问题答题库", 1)
    add_heading(doc, "AI共创界面", 2)
    add_question(doc, "这个项目想解决的到底是什么", "当用户希望和AI一起创作有个人特色的作品时，模糊输入容易被AI按大众语义理解，结果变得普通。项目让用户看见AI如何解释模糊词，并决定是否补充自己的标准。")
    add_question(doc, "你的核心设计动作是什么", "在AI对模糊词的解释旁提供确认和调整入口，通过按钮引导用户判断“这是不是我的意思”。它把原本隐性的价值判断变成用户可以参与的步骤。")
    add_question(doc, "研究怎么做", "前期通过AI对话文本和主观评价分析理解分歧及原因；后续计划比较界面方案，并用眼动观察用户在真实输入时对解释区和提示区的注意。", "不要说已经完成眼动实验。")

    add_heading(doc, "好久没吃小程序", 2)
    add_question(doc, "它和普通收藏夹有什么不同", "它不是记录吃过什么，而是帮助用户为不同食物设定自己的参考周期。用户想吃点好久没吃的时，可以结合距离上次食用的天数判断现在是否值得再吃。")
    add_question(doc, "为什么要做成真实小程序", "概念问卷只能判断用户是否理解需求。做成可用小程序后，可以检查记录、查询和周期设置的流程是否完整，也能证明我能用AI开发工具把产品概念落到可使用的原型。")

    add_heading(doc, "FrameTrace", 2)
    add_question(doc, "为什么摄影需要MR", "人在陌生或值得拍摄的场景里，常在入口拍一张全景就结束，较少继续寻找机位、构图和焦段。MR可以把其他拍摄者的机位、方向和姿势作为空间线索，让用户在真实场景中继续探索。")
    add_question(doc, "问卷结果说明了什么", "前期用概念模拟示意图收集106份问卷，排除高风险回答后以82份作为主分析，用于了解第一印象、理解和使用顾虑。它说明概念的可理解性与担忧，不等于已经验证真实拍摄行为会改变。")

    add_heading(doc, "札幌市路面电车广告价值提升", 2)
    add_question(doc, "这个项目和产品岗位有什么关系", "它是一个面向真实限制条件的方案设计：新旧电车同时运行，旧车广告更自由、更受欢迎。我们提出用同一企业的过去和未来连接两类车，并用AI工作流批量研究企业特征、生成符合车身广告限制的案例，再进行概念调研。")
    add_question(doc, "你在AI工作流中做了什么", "我不是只让AI出图，而是把企业历史、未来形象、标志性特征和车身广告限制组织成可复现的步骤，再人工检查案例是否符合表达和限制条件。")

    add_heading(doc, "高频面试问题", 1)
    add_question(doc, "你为什么从日本回国求职", "我在日本的学习让我更重视用户感受和研究方法，但我希望回到国内参与AI、智能硬件和智能汽车这些变化很快、产品落地规模更大的场景。回国后我想把研究方法放进真实产品迭代中。")
    add_question(doc, "你的专业和国内岗位名称有什么对应关系", "本科是工业设计，硕士阶段是人间情报设计，核心关注人机交互和用户体验研究。我的能力不是只做视觉界面，而是从使用场景、主观感受和行为线索出发，把问题整理成交互方案和可验证原型。")
    add_question(doc, "你没有正式实习，凭什么能做产品工作", "我的项目有完整的场景、调研、方案和验证材料，其中好久没吃已经做成可用小程序，冰箱也做过带电控的实体模型。我的短板是缺少企业协作经验，所以我希望从产品定义、用户研究和原型验证这些能快速承担的工作开始，并主动补足研发协作和业务节奏。", "不要回避实习经历空白，也不要把学校项目说成公司项目。")
    add_question(doc, "你会哪些工具", "Figma用于流程和交互原型，Blender用于产品场景和空间表达，Photoshop和Illustrator用于视觉材料，After Effects用于操作演示。Codex等AI工具用于原型开发和资料整理。工具是为了验证概念，不是为了把自己包装成纯执行型软件岗位。")
    add_question(doc, "你怎么和工程师合作", "我会先把用户场景、关键流程、输入输出和异常情况整理清楚，再用原型和说明把设计取舍讲明白。遇到技术限制时，先确认限制影响的是功能目标还是表现方式，再一起找能保留核心体验的实现路径。")
    add_question(doc, "如果面试官质疑你的项目数据不够", "我会直接说明目前数据的用途和边界：概念调研用于判断用户是否理解和愿意尝试，不能证明长期效果。数据不够时，我会提出下一轮需要补什么，例如任务测试、对比原型、有效样本筛选或更接近真实场景的观察。")

    add_heading(doc, "三种现场自我介绍", 1)
    add_heading(doc, "科技产品版", 2)
    add_body(doc, "我是孔维鹏，本科是工业设计，目前在日本札幌市立大学读人间情报设计硕士，方向是人机交互和产品体验。我主要做AI产品体验、用户研究和交互原型：研究用户怎样参与AI的关键判断，也用AI开发工具做过上线的微信小程序。今天想了解贵司产品体验、AI产品或用户研究相关的2027届岗位。")
    add_heading(doc, "智能硬件版", 2)
    add_body(doc, "我是孔维鹏，本科工业设计，目前在日本读人间情报设计硕士。我关注智能硬件在真实生活场景中怎样通过设备、灯光和界面给用户清晰反馈。我的毕业项目做过冰箱内食材日期与位置管理，完成了场景模拟和带电控的1/4模型。今天想了解贵司智能硬件、产品体验或产品设计方向的岗位。")
    add_heading(doc, "智能汽车HMI版", 2)
    add_body(doc, "我是孔维鹏，本科工业设计，目前在日本读人机交互方向硕士。我关注自动驾驶情境下，系统如何让驾驶员理解风险来源，同时维持对道路和驾驶的关注。我做过智能驾驶HMI概念研究，结合问卷和后续眼动方案评估提示方式。今天想了解贵司智能座舱、HMI或产品体验方向的岗位。")

    add_heading(doc, "招聘会当天的准备清单", 1)
    add_bullet(doc, "手机中离线保存三份简历和作品集PDF。现场先确认部门与岗位，再递对应版本。")
    add_bullet(doc, "准备一份可快速打开的项目目录：海尔看冰箱，汽车企业看HMI，腾讯和荣耀看AI共创或小程序。")
    add_bullet(doc, "每次交流先用20到30秒说明背景和目标，再问岗位。对方感兴趣后才展开一个项目，不从作品集封面开始翻。")
    add_bullet(doc, "主动记录企业、岗位、HR姓名或联系方式，以及对方追问的项目点。当天结束后按记录补投电子简历或发送作品集链接。")
    add_bullet(doc, "遇到不知道的问题，不要硬答。可以说“我目前没有做过这一部分，但我理解它需要先确认……，我会从……开始补”。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
