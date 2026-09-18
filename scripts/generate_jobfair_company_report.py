from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "research"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "2026-09-19-shenzhen-jobfair-company-research-cn.docx"

BLACK = "000000"
NAVY = "17365D"
BLUE = "2F5D8A"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F6F8"
MID_GRAY = "687383"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial Unicode MS", size=10.5, bold=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "19")
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_para(doc: Document, text: str = "", style: str | None = None, *, bold_lead: str | None = None,
             space_after=5, line_spacing=1.14, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if keep:
        keep_with_next(p)
    return p


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, size=10.2)
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.1


def set_col_width(cell, width_inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_table(table, widths=None, header=True, font_size=9.2):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if header:
        set_repeat_table_header(table.rows[0])
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths and j < len(widths):
                set_col_width(cell, widths[j])
            if i == 0 and header:
                set_cell_shading(cell, NAVY)
            elif i % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, WHITE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(i == 0 and header), color=(WHITE if i == 0 and header else BLACK))


def add_two_col_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.add_row().cells[0].text = "项目"
    table.rows[0].cells[1].text = "调研结果"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    style_table(table, widths=[1.48, 5.42], header=True, font_size=9.35)
    for row in table.rows[1:]:
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_source_line(doc: Document, sources: list[tuple[str, str]]):
    p = doc.add_paragraph()
    lead = p.add_run("资料来源  ")
    set_run_font(lead, size=9.2, bold=True, color=MID_GRAY)
    for idx, (label, url) in enumerate(sources):
        if idx:
            r = p.add_run("  |  ")
            set_run_font(r, size=9.2, color=MID_GRAY)
        add_hyperlink(p, label, url)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end = paragraph.add_run(" 页")
    set_run_font(end, size=8.5, color=MID_GRAY)


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


companies = [
    {
        "name": "德赛西威",
        "priority": "S",
        "match": "高度匹配",
        "status": "业务扩张  2027校招进行中",
        "overview": "惠州市德赛西威汽车电子股份有限公司是汽车一级供应商和上市公司，业务集中在智能座舱、智能驾驶和网联服务。客户面对的是整车厂，工作通常要同时处理用户体验、车规安全、软硬件约束和量产交付。",
        "products": [
            "智能座舱  座舱域控制器、仪表、信息娱乐系统、车载显示、HUD、智能交互和音频产品",
            "智能驾驶  高算力域控制器、摄像头、毫米波雷达、泊车及驾驶员监测",
            "软件与连接  车联网、生态软件、软件平台和数字钥匙",
            "近期方向  端侧大模型、多模态交互、3D HMI、主动感知和舱驾融合",
        ],
        "current": "2025年营业收入约325.6亿元。2026年上半年披露的第五代智能座舱平台已在理想汽车规模化量产；公司仍在建设新园区并扩大全球研产网络。公司业务与智能汽车增长周期联系紧密，客户项目、车规验证和量产节点会直接影响岗位节奏。",
        "event_jobs": "活动公开岗位以大模型、AI智能体、模型部署、AI应用、定位、地图和底盘算法为主。公司独立发布的2027校招还设有产品经理和项目经理岗位，工作地点包括惠州、深圳、上海、南京和成都。",
        "role_read": "产品经理可能负责平台规划、客户需求、功能定义与方案落地；项目经理可能更偏进度、质量、客户接口和量产交付。HMI和用户研究岗位未在活动清单中单独列名，需要向展位人员确认内部团队和岗位编号。",
        "fit": "智能驾驶HMI、驾驶注意与风险表达项目可以直接对应座舱体验；眼动、主观评价和场景分析可对应体验测试；工业设计背景有助于理解显示、交互和硬件约束。",
        "goals": [
            "确认产品项目类岗位中是否有智能座舱HMI、用户体验研究或交互验证方向",
            "确认岗位产出是场景与功能定义、HMI方案，还是客户交付和项目排期",
            "确认HCI和工业设计背景是否满足专业要求，以及2027年4月海外毕业的届别认定",
            "争取拿到具体事业部、岗位编号和作品集投递联系人",
        ],
        "risks": "岗位总体偏工程技术。若现场只能提供算法和软件研发岗位，则不要强行投递；可以留下针对座舱产品和体验研究的简历，并询问后续产品项目岗位。",
        "sources": [
            ("公司解决方案", "https://www.desaysv.com/solution.html"),
            ("2026半年报", "https://disc.static.szse.cn/disc/disk03/finalpage/2026-08-14/2aab2da4-0d48-471f-b396-fbded887d69c.PDF"),
            ("2027招聘简章", "https://myjob.dlmu.edu.cn/campus/view/id/868670"),
        ],
    },
    {
        "name": "荣耀",
        "priority": "S",
        "match": "高度匹配",
        "status": "AI终端转型  2027校招进行中",
        "overview": "荣耀是总部位于深圳的AI终端公司，产品覆盖手机、平板、笔记本、穿戴、音频和AIoT。公司正在从单一终端厂商转向系统级AI和跨设备生态，重点包括MagicOS、YOYO智能体、端侧AI和设备协同。",
        "products": [
            "移动终端  手机、折叠屏、平板、笔记本和可穿戴设备",
            "系统体验  MagicOS、YOYO智能体、意图识别和跨设备服务",
            "AI生态  HONOR AI Connect、端云协同和多品牌设备连接",
            "新形态  机器人手机和系统级Agent架构",
        ],
        "current": "官方披露员工超过1.4万人，研发人员占比超过70%，研发投入占营收比例约11.5%。2026年继续发布AI终端和机器人手机，2027届校园招聘已启动。公司是非上市企业，外部无法像上市公司一样获得完整财务信息，因此应把产品进展、招聘节奏和团队信息作为主要判断依据。",
        "event_jobs": "产品与运营经理、通用软件开发、硬件开发、销售解决方案培训生、营销创意管培生和服务解决方案培训生。",
        "role_read": "产品与运营经理可能覆盖系统产品、应用产品、海外运营、用户增长或渠道运营，岗位名称不足以判断是否属于AI体验。服务解决方案岗位也可能偏售后服务设计和渠道支持。",
        "fit": "AI共创界面、个人判断参与机制、跨设备智能硬件和用户研究可对应系统级AI体验。日语和海外学习经历也可支持海外产品、本地化体验和日本市场方向。",
        "goals": [
            "确认产品与运营经理属于MagicOS、YOYO、AI终端、IoT还是营销运营",
            "询问是否存在AI交互、用户研究、跨设备体验或系统产品岗位",
            "确认是否需要作品集以及产品岗位更重视研究、原型还是数据运营",
            "询问日本市场或国际产品团队是否需要日语和海外用户研究能力",
        ],
        "risks": "岗位竞争强，产品岗位可能优先计算机、通信或有成熟实习经历的候选人。回答时需要把研究项目讲成可落地的产品问题、功能决策和验证方法。",
        "sources": [
            ("荣耀简介", "https://www.honor.com/cn/brand/"),
            ("荣耀招聘", "https://www.honor.com/cn/career/"),
            ("AI生态进展", "https://www.honor.com/cn/news/honor-ai-ecosystem/"),
        ],
    },
    {
        "name": "海尔智家",
        "priority": "S",
        "match": "高度匹配",
        "status": "经营成熟  全面推进家庭AI",
        "overview": "海尔智家是全球化家电与智慧家庭企业，业务已从冰箱、洗衣机、空调等单品扩展到UHomeOS、智家App、家庭AI智能体、机器人和人车家联动。公司同时拥有海尔、卡萨帝、统帅以及多个海外品牌。",
        "products": [
            "家电产品  冰箱、洗衣机、空调、厨电、热水器和清洁设备",
            "智慧家庭  智家App、UHomeOS、3D家庭视图和全屋场景联动",
            "AI产品  小优Agent、食材识别、主动服务和家庭空间知识图谱",
            "全球平台  国内HaiSmart及欧美澳SmartHQ等本地化平台",
        ],
        "current": "公司在A股、D股和H股上市。2026年仍在推进AI产品、AI场景和AI生活，官方披露智家App月活超过1600万。2027届招聘已启动，并公开AI产品经理、产品设计师、UI设计师和CMF设计师等岗位。",
        "event_jobs": "硬件研发、软件研发、AI与算法、智能制造、智能类、设计类、智慧物流和全球人才发展项目。设计和AI岗位在官方校招中有更细分的信息。",
        "role_read": "AI产品经理可能负责家庭智能体、数据服务或企业数字化产品；设计类可能进入单品产品设计、UI、CMF或场景体验。岗位多集中在青岛，广东岗位需要现场确认。",
        "fit": "冰箱物品日期与位置管理项目能对应食材管理和智慧厨房；饮食周期小程序对应生活服务与行为设计；AI交互研究可对应家庭智能体和主动服务。",
        "goals": [
            "确认AI产品经理、产品设计师和UI设计师的具体团队与工作地点",
            "询问家庭AI Agent是否招聘交互、用户研究或场景定义人才",
            "了解广东地区是否有智慧家庭、产品体验或研发团队",
            "确认产品设计岗位是否接受包含服务系统和交互研究的作品集",
        ],
        "risks": "若不能接受青岛工作，需要尽早筛选地域。部分岗位会按产业和品类分配，入职后的实际产品方向可能与投递名称不完全一致。",
        "sources": [
            ("2026 AI进展", "https://www.haier.com/press-events/news/20260721_293366.shtml"),
            ("2026半年报", "https://static.cninfo.com.cn/finalpage/2026-08-28/1225522691.PDF"),
            ("海尔校招", "https://hd.maker.haier.net/client/campus/index"),
        ],
    },
    {
        "name": "腾讯",
        "priority": "A",
        "match": "方向匹配但岗位不明确",
        "status": "经营稳定  AI投入与校招持续",
        "overview": "腾讯业务覆盖微信、QQ、游戏、内容、广告、金融科技、腾讯云和AI。公司体量大，产品和设计岗位分散在不同事业群，企业名称本身不足以判断实际工作内容。",
        "products": [
            "消费互联网  微信、QQ、内容与数字娱乐",
            "企业服务  腾讯云、会议、文档、营销和行业解决方案",
            "人工智能  混元模型、元宝、智能体与多媒体实验室技术",
            "出行相关  地图、车联、云服务及面向汽车行业的解决方案",
        ],
        "current": "腾讯已发布2026年第二季度业绩并启动2027全球校园招聘。海外校招面向2026年1月至2027年12月毕业的学生。公开招聘仍覆盖技术、产品、设计、市场和职能岗位。",
        "event_jobs": "活动只公布技术、产品、设计、市场和职能五个大类，没有事业群、具体岗位和工作地点。",
        "role_read": "产品岗位可能是C端产品、商业产品、AI产品、游戏策划或企业服务；设计岗位可能是交互、视觉、用户研究、游戏设计或品牌设计。必须先确认岗位编号和业务团队。",
        "fit": "AI产品体验、用户判断机制、小程序开发和用户研究可以对应AI产品、C端体验和用户研究。车载HMI只有在地图、车联或汽车云团队才直接相关。",
        "goals": [
            "获得具体事业群、岗位名称和岗位编号",
            "确认现场投递是否会占用官网校招的投递机会",
            "优先寻找AI产品、用户研究、交互设计、智能硬件或地图出行岗位",
            "询问作品集评估方式以及研究型项目如何进入产品面试流程",
        ],
        "risks": "若展位只能统一收简历且不能说明业务团队，现场信息价值有限。不要因为公司知名度而投递一个方向模糊的岗位。",
        "sources": [
            ("2026业绩页面", "https://www.tencent.com/zh-cn/investors/results/"),
            ("腾讯招聘", "https://hr.tencent.com/zh-cn/jobopportunity.html"),
            ("海外校招说明", "https://careers.tencent.com/m/zh-cn/faqdetail.html"),
        ],
    },
    {
        "name": "吉利",
        "priority": "A",
        "match": "汽车方向匹配",
        "status": "销量与收入增长  智能化招聘活跃",
        "overview": "吉利控股旗下包含吉利、银河、领克、极氪等汽车品牌，并布局智能座舱、智能驾驶、车载AI、能源和全球化研发。整车企业的产品岗位通常更接近车型定义和用户体验，但也会受到车型周期、成本和供应链约束。",
        "products": [
            "整车品牌  燃油、混动和纯电车型及多个子品牌",
            "智能座舱  Flyme Auto、座舱芯片、语音、车载AI和多屏交互",
            "智能驾驶  感知、规控、底盘域控和驾驶辅助",
            "全球业务  海外车型、区域适应性开发和多地设计研发中心",
        ],
        "current": "2026年上半年吉利汽车收入约1736亿元，创同期纪录；核心利润同比增长46%。2027届全球校招开放1000多个岗位，覆盖35个以上城市，包含AI智能座舱、端侧多模态、车载大模型和产品管理。",
        "event_jobs": "人工智能算法、大模型算法、智能驾驶软件、新能源电源、智能声学、产品管理、海外项目管理和全球适应性开发。",
        "role_read": "产品管理可能是车型产品、智能座舱、平台规划或生命周期管理；全球适应性开发可能涉及法规、环境、用户习惯和本地化测试，不一定是纯用户研究。",
        "fit": "车载HMI项目直接对应智能座舱；日语N1和日本学习经历可支持日本市场、本地化体验和全球适应性开发；工业设计背景有利于连接交互、硬件与整车场景。",
        "goals": [
            "确认产品管理岗位是否包含智能座舱HMI或用户体验方向",
            "询问全球适应性开发是否需要日本用户研究、语言和本地化能力",
            "确认工作地点、品牌或事业部，以及岗位是否面向2027届海外毕业生",
            "了解作品集在产品管理岗位中的评估权重",
        ],
        "risks": "岗位可能集中在杭州、宁波、上海等地。若产品管理实际偏车型成本和项目节点，需要判断自己是否接受较强的工程与业务协调属性。",
        "sources": [
            ("2026中期业绩", "https://newsroom.geely.com/geely-auto-record-h1-2026-revenue"),
            ("吉利汽车投资者信息", "https://www.geelyauto.com.hk/zh-cn/%E5%85%AC%E5%8F%B8%E7%AE%80%E4%BB%8B/"),
            ("2027招聘信息", "https://career.hebut.edu.cn/home/correcruit/content/id/79948.html"),
        ],
    },
    {
        "name": "蓝思科技",
        "priority": "A减",
        "match": "智能硬件与设计匹配",
        "status": "新业务转型  短期业绩承压",
        "overview": "蓝思科技从玻璃、金属和精密结构件制造扩展到消费电子整机、智能座舱、AI眼镜、XR、人形机器人和AI服务器。其优势是材料、工艺、自动化和大规模量产，产品岗位通常与制造可行性和客户交付联系紧密。",
        "products": [
            "AI智能终端  手机、PC、穿戴、AI眼镜与XR相关结构件和整机服务",
            "智能汽车  座舱交互系统、显示与结构件",
            "具身智能  机器人核心部件、结构件与整机制造",
            "制造能力  新材料、自动化设备、模组、软件开发和整机组装",
        ],
        "current": "2026年上半年营业收入288.66亿元，同比下降12.42%；归母净利润5.77亿元，同比下降49.52%。二季度利润同比略增并环比恢复。公司在传统消费电子需求承压时继续投入AI硬件、机器人和智能座舱。",
        "event_jobs": "ID设计工程师、研发项目经理、产品市场专员、项目专员、结构设计工程师、基带、驱动、系统软件、IoT软件及职能岗位。",
        "role_read": "ID设计可能强调材料、曲面、CMF和量产；研发项目经理可能负责客户需求、开发节点、成本和交付；产品市场更接近行业研究、客户提案和产品路线。",
        "fit": "工业设计和智能硬件项目适合ID与研发项目方向；HCI项目能补充用户洞察，但必须证明自己理解结构、材料和量产限制。",
        "goals": [
            "确认岗位属于消费电子、汽车座舱、机器人还是传统结构件事业部",
            "确认ID设计负责整机、交互硬件还是零部件造型",
            "了解研发项目经理的产品定义与量产交付占比",
            "确认工作地点、驻厂比例和项目加班节奏",
        ],
        "risks": "公司经营规模大，但2026上半年业绩承压。不同事业部差异很大，不能只看公司名称，需要把岗位、地点和产品线一起判断。",
        "sources": [
            ("2026半年报摘要", "https://static.cninfo.com.cn/finalpage/2026-08-22/1225491664.PDF"),
            ("2025年报", "https://disc.static.szse.cn/disc/disk03/finalpage/2026-03-31/51a36b51-bc87-410b-8fca-aaab38dacbd9.pdf"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
    {
        "name": "锐明技术",
        "priority": "B加",
        "match": "商用车产品备选",
        "status": "收入增长  利润下降  AI产品持续投入",
        "overview": "锐明技术为商用车提供AI视频监控、驾驶员监测、ADAS、车队管理和行业解决方案，客户包括公交、物流、出租、校车和特种车辆运营者。其产品价值更强调安全、运营效率和法规合规。",
        "products": [
            "车载设备  专业摄像机、录像与计算终端、传感器和显示设备",
            "安全产品  驾驶员疲劳监测、ADAS、事件识别和风险预警",
            "平台服务  车队管理、远程监控、数据分析和FT Cloud",
            "AI Agent  路径规划、故障预诊断和运营决策辅助",
        ],
        "current": "2026年上半年营收约12.38亿元，同比增长7.09%；归母净利润约1.11亿元，同比下降45.33%。公司继续增加AI感知、决策和应用投入，并拓展日本、印度等海外市场。",
        "event_jobs": "海外技术支持、海外销售、产品工程师、产品市场工程师、硬件工程师和自动驾驶感知算法工程师。",
        "role_read": "产品工程师可能负责规格、测试、客户需求和解决方案落地；产品市场工程师可能负责行业研究、竞品、销售材料和客户方案。两者都可能比消费互联网产品经理更技术化。",
        "fit": "驾驶注意、风险提示和HMI研究可对应驾驶员监测与预警产品；日语可支持日本市场；但商用车产品强调B端运营和法规，需要转换作品集叙事。",
        "goals": [
            "确认产品工程师是否参与需求分析、产品定义和车载终端交互",
            "询问驾驶员监测与预警产品是否做HMI或用户测试",
            "了解日本市场团队是否需要日语与本地化研究能力",
            "确认专业限制和技术知识要求",
        ],
        "risks": "盈利同比下降需要持续观察，但公司仍盈利并保持收入增长。岗位可能偏技术支持和交付，需避免把产品工程师误认为体验产品经理。",
        "sources": [
            ("2026半年报信息", "https://stock.stockstar.com/notice/SN2026082600050138.shtml"),
            ("公司官网", "https://www.streamax.com/"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
    {
        "name": "视源股份",
        "priority": "B加",
        "match": "智能交互硬件匹配",
        "status": "企业服务增长  AI与机器人投入增加",
        "overview": "视源股份总部位于广州，核心品牌包括教育场景的希沃和会议场景的MAXHUB，同时经营显示控制板、交互智能平板、工业计算和机器人业务。公司擅长把显示硬件、软件和垂直场景结合。",
        "products": [
            "教育产品  希沃交互平板、教学软件、AI教育解决方案",
            "企业服务  MAXHUB会议平板、协作软件和行业解决方案",
            "工业产品  工业主板、边缘计算、工控整机和机器人计算板卡",
            "机器人  商用清洁机器人及其他场景化设备",
        ],
        "current": "2026年上半年企业服务品牌业务收入17.14亿元，同比增长38.43%；MAXHUB海外收入4.38亿元，同比增长50.42%。公司持续把AI能力部署到教育、会议、医疗、金融、工业和机器人终端。",
        "event_jobs": "应用软件、嵌入式系统、软件技术支持、数据分析、硬件、电源、算法和销售。活动清单未单列产品经理或设计岗位。",
        "role_read": "现场岗位与HCI并非直接对应，但公司长期需要产品、交互和用户研究人才。展位的主要价值是确认是否有未公开岗位或能否转交希沃、MAXHUB产品团队。",
        "fit": "智能硬件、跨设备交互、AI体验和用户研究适合会议及教育产品；广州工作地点也具有现实优势。",
        "goals": [
            "询问是否有未列出的产品经理、交互设计或用户研究岗位",
            "确认希沃、MAXHUB和机器人团队的2027届招聘需求",
            "了解作品集是否可以直接转交产品体验团队",
            "判断技术支持岗位是否包含产品反馈闭环，还是纯客户服务",
        ],
        "risks": "不要为了进入公司而投递明显不匹配的软件或硬件研发岗位。若没有产品体验机会，可保留联系方式等待后续岗位。",
        "sources": [
            ("2026半年报", "https://static.cninfo.com.cn/finalpage/2026-08-27/1225511876.PDF"),
            ("视源官网", "https://www.cvte.com/"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
    {
        "name": "快递100",
        "priority": "B加",
        "match": "AI产品经理备选",
        "status": "私营科技企业  AI产品化活跃",
        "overview": "快递100不是快递运输公司，而是连接快递物流企业、消费者、快递员和企业客户的信息平台。其公司主体为深圳前海百递网络有限公司，早期由金蝶孵化。",
        "products": [
            "消费者产品  快递查询、寄件App、小程序、快应用和智能体",
            "从业者产品  快递100收件端和网点经营工具",
            "企业产品  百递云API、企业快递SaaS和电商快递管理",
            "AI能力  物流网络数智图谱、大模型应用平台和AI快递助手",
        ],
        "current": "公司公开信息显示其服务个人、快递从业者和企业三类用户，正推进AI IN ALL战略，把AI能力嵌入查询、寄件、管理和客服。作为非上市企业，公开财务信息有限，应重点确认产品规模、团队和岗位责任。",
        "event_jobs": "AI研究员、数据工程师、产品经理、前后端、测试、产品运营、产品市场、国际业务和技术支持。",
        "role_read": "产品经理可能分为C端、B端SaaS、物流API和AI Agent。不同方向对用户研究、业务流程、数据分析和技术理解的要求差异很大。",
        "fit": "AI共创研究、功能定义、小程序开发和用户研究均可转化为AI产品经理案例。它不能延续车载HMI路线，但适合建立真实上线和业务闭环经验。",
        "goals": [
            "确认产品经理负责C端、B端SaaS、API还是AI Agent",
            "了解新人是否能承担需求研究、原型、上线和数据复盘",
            "询问AI产品岗位需要的技术深度和面试作业",
            "确认团队规模、汇报对象和试用期考核指标",
        ],
        "risks": "公司对外宣传信息丰富，但非上市公司财务透明度有限。需要通过岗位职责、团队人数、产品指标和面试流程判断机会质量。",
        "sources": [
            ("公司介绍", "https://www.kuaidi100.com/about/index.shtml"),
            ("产品服务", "https://m.kuaidi100.com/about/index.jsp"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
    {
        "name": "黑白调和傲风",
        "priority": "B",
        "match": "工业设计与用户体验匹配",
        "status": "2027校招进行中  智能化方向开始扩展",
        "overview": "黑白调和傲风由浙江分享空间科技有限公司运营。黑白调聚焦人体工学办公椅、儿童桌椅和企业健康办公；傲风聚焦电竞椅、电竞桌和电竞舱。产品核心仍是家具与人体工学，但招聘中已出现嵌入式、算法和AI部署岗位。",
        "products": [
            "人体工学产品  办公椅、升降桌和儿童学习桌椅",
            "电竞装备  电竞椅、电竞桌、电竞舱及赛事合作产品",
            "智能化方向  嵌入式软硬件、算法、AI前沿部署和智能坐姿相关可能性",
            "全球业务  跨境电商、海外营销和多语言市场运营",
        ],
        "current": "公司2027校招面向2026年9月至2027年8月毕业的海内外学生，工作地点为杭州和深圳。公开岗位覆盖产品研发、用户体验、设计、跨境运营和市场。",
        "event_jobs": "ID设计、CMF设计、产品工程、结构、嵌入式软硬件、算法、用户体验、产品营销、3D渲染和跨境岗位。",
        "role_read": "用户体验岗位可能是前期研究，也可能偏售后问题与满意度改善；产品工程师偏结构、打样和量产；ID与CMF更直接匹配工业设计背景。",
        "fit": "工业设计和产品服务系统能力匹配；用户研究与人体工学可以形成特色；AI/HCI研究只有在智能坐姿或软件服务产品中才能充分使用。",
        "goals": [
            "确认用户体验岗位是前期研究还是售后体验改善",
            "询问是否有智能坐姿检测、传感器或App协同产品",
            "确认ID和CMF岗位的作品集要求、设计测试与工作地点",
            "了解深圳团队承担研发、运营还是海外业务",
        ],
        "risks": "不要把品牌宣传中的智能化方向等同于成熟AI产品线。应通过实际产品、团队人数和岗位任务判断智能业务是否已经落地。",
        "sources": [
            ("2027招聘信息", "https://job.henu.edu.cn/module/position_brief_detail/id-113762/nid-10822"),
            ("黑白调官网", "https://www.heibaidiao.com/WalkInto"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
    {
        "name": "ZURU",
        "priority": "B减",
        "match": "国际消费品设计备选",
        "status": "全球化私营企业  招聘流程成熟",
        "overview": "ZURU是国际消费品集团，业务包含玩具、快速消费品和建筑科技。深圳是其产品设计、工程、采购和供应链的重要基地。公司强调快速开发、全球市场和跨文化协作。",
        "products": [
            "ZURU Toys  玩具、儿童娱乐产品和全球零售品牌",
            "ZURU Edge  快速消费品和日用产品",
            "ZURU Tech  建筑与住宅科技相关业务",
            "中国团队  设计、研发、制造、采购和供应链协同",
        ],
        "current": "官方招聘页面显示公司在全球30多个地点设有团队，深圳设有办公室。招聘流程通常包含初步沟通、直属负责人面试、岗位作业与展示，以及多轮团队面试。",
        "event_jobs": "新品研发项目管理管培生、初级3D设计师、初级Free Form设计师、平面设计、包装工程、玩具工程和供应链管培生。",
        "role_read": "3D和Free Form岗位更偏造型、曲面与产品实现；新品研发项目管理可能参与概念、成本、供应商和上市进度；用户研究和HCI不是岗位核心。",
        "fit": "工业设计、Blender和海外协作能力匹配。若希望继续AI产品或车载HMI路线，ZURU的长期方向偏离较大。",
        "goals": [
            "确认3D与Free Form岗位的具体软件、产品类型和工程深度",
            "了解新品研发项目岗位是否参与用户洞察与产品定义",
            "询问岗位作业形式、完成时间和评价标准",
            "确认工作地点、产品事业部和英语使用频率",
        ],
        "risks": "岗位可能要求高强度快速迭代和较强造型能力。需要判断自己的作品集是否有足够的实体产品、曲面和量产表达。",
        "sources": [
            ("全球招聘", "https://zuru.com/careers/"),
            ("中国校招", "https://www.zururecruit.cn/campus-recruitment"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
    {
        "name": "爱奇迹",
        "priority": "C",
        "match": "岗位匹配但行业需谨慎",
        "status": "海外业务导向  公开财务信息有限",
        "overview": "公开招聘资料显示，爱奇迹的核心业务是电子雾化产品，主要面向海外市场。公司长期招聘产品、工业设计、CMF、UI、结构、电子、材料和海外营销岗位。",
        "products": [
            "电子雾化硬件  面向海外市场的雾化设备及相关配件",
            "产品研发  外观、CMF、结构、电子、材料、测试和量产",
            "海外业务  GTM、渠道销售、零售、数字营销和本地化",
            "设计与品牌  工业设计、UI、创意设计和海外品牌传播",
        ],
        "current": "公司为非上市企业，公开财务和客户信息有限。既往校招显示总部岗位位于深圳，产品经理职责涉及用户洞察、产品规划、产品定义和生命周期管理，并偏好有海外经历及电子硬件知识的候选人。",
        "event_jobs": "海外营销、产品经理、CMF与UI、材料、结构、电子、项目管理、供应链和职能岗位。",
        "role_read": "产品经理可能承担海外用户洞察、产品定义、供应链协调和上市节奏。CMF与工业设计较直接，但行业监管与出口合规会显著影响产品策略。",
        "fit": "海外学习、日语、工业设计和用户研究均有价值；电子硬件知识不足可能成为产品经理岗位门槛。",
        "goals": [
            "先确认自己是否接受电子雾化行业，再决定是否投递",
            "确认主要市场、具体产品线、岗位合同主体和工作地点",
            "询问产品经理的用户研究、产品定义和供应链工作比例",
            "了解生产与出口资质、合规团队及日本市场需求",
        ],
        "risks": "行业存在健康、监管和个人价值观问题。对公司信息透明度、合规资质和合同主体的核实应高于一般消费电子企业。",
        "sources": [
            ("招聘岗位资料", "https://gdhr.bibibi.net/detail/job?id=2723186"),
            ("历史校园宣讲", "https://career.gdufs.edu.cn/web/index/preach-detail?id=567"),
            ("活动岗位清单", "https://www.sohu.com/a/1074867336_121106875"),
        ],
    },
]


def build_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = "Arial Unicode MS"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 17, 15, 7), ("Heading 2", 13, 11, 5), ("Heading 3", 11, 8, 3)):
        st = styles[name]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    footer = sec.footer
    fp = footer.paragraphs[0]
    left = fp.add_run("9月19日深圳海归招聘会企业调研")
    set_run_font(left, size=8.5, color=MID_GRAY)
    fp.add_run(" " * 12)
    add_page_number(fp)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("9月19日深圳海归招聘会企业调研报告")
    sub = doc.add_paragraph()
    r = sub.add_run("面向AI产品体验  智能硬件  车载HMI与用户研究岗位")
    set_run_font(r, size=14, bold=True, color=NAVY)
    sub.paragraph_format.space_after = Pt(28)

    meta = doc.add_table(rows=5, cols=2)
    meta_data = [
        ("候选人", "孔维鹏"),
        ("教育背景", "工业设计本科  日本人机交互硕士在读  预计2027年4月毕业"),
        ("研究范围", "深圳海归人才招聘会已公开企业与岗位"),
        ("研究基准日", "2026年9月18日"),
        ("用途", "现场筛选企业  识别岗位  准备提问  记录后续行动"),
    ]
    for row, data in zip(meta.rows, meta_data):
        row.cells[0].text, row.cells[1].text = data
    style_table(meta, widths=[1.38, 5.52], header=False, font_size=10)
    for row in meta.rows:
        set_cell_shading(row.cells[0], PALE_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    intro = add_para(doc, "本报告用于支持9月19日现场沟通和后续投递。优先结论是：德赛西威、荣耀和海尔智家最接近现有作品集；腾讯、吉利和蓝思科技值得投入，但必须先确认具体团队；视源股份、锐明技术、快递100和黑白调可作为清晰的第二梯队。企业公开名单和岗位可能继续更新，最终信息以展位和正式招聘系统为准。", space_after=8, line_spacing=1.25)
    intro.paragraph_format.first_line_indent = Inches(0.24)
    p = doc.add_paragraph()
    add_hyperlink(p, "招聘会最新公开企业与岗位清单", "https://www.sohu.com/a/1074867336_121106875")

    doc.add_page_break()
    add_heading(doc, "报告使用方法", 1)
    add_para(doc, "每家公司按照业务、产品、发展状态、招聘岗位、岗位含义、个人匹配、现场目标和风险八个维度整理。经营状态主要使用公司官网、交易所公告、定期报告和招聘官网；非上市公司的财务透明度较低，因此不将品牌宣传直接等同于经营实力。")
    add_heading(doc, "现场判断标准", 2)
    add_bullets(doc, [
        "岗位身份  确认正式岗位名称、岗位编号、招聘届别和劳动合同主体",
        "工作内容  区分产品定义、体验研究、项目交付、技术支持和销售运营",
        "团队位置  确认事业部、产品线、工作城市、汇报对象和团队人数",
        "能力证据  确认简历、作品集、设计测试、笔试和面试所需材料",
        "后续动作  记录投递链接、联系人、截止日期和下一轮时间",
    ])

    add_heading(doc, "候选人优势与需要补强的证据", 2)
    add_two_col_table(doc, [
        ("可直接使用的优势", "工业设计基础  HCI研究  日本学习经历  日语N1  AI交互  车载HMI  用户研究  眼动与主观评价  原型和小程序开发"),
        ("产品岗位需要补强", "把研究过程表达成需求判断、功能定义、取舍、验证结果和后续迭代，而不是只展示概念与视觉"),
        ("硬件岗位需要补强", "结构、材料、制造、成本、技术可行性与跨团队协作证据"),
        ("汽车岗位需要补强", "车规安全、驾驶场景、信息优先级、法规限制和量产流程的基本理解"),
    ])

    add_heading(doc, "企业优先级总览", 1)
    table = doc.add_table(rows=1, cols=5)
    headers = ["企业", "优先级", "主要机会", "经营或业务状态", "现场策略"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    summary_rows = [
        ("德赛西威", "S", "智能座舱产品  HMI  体验研究", "扩张与量产", "必须问到具体团队"),
        ("荣耀", "S", "AI产品  系统体验  跨设备交互", "AI终端转型", "区分产品与运营"),
        ("海尔智家", "S", "AI产品  智慧家庭  产品设计", "成熟且推进AI", "确认工作城市"),
        ("腾讯", "A", "AI产品  用户研究  交互设计", "稳定  岗位宽泛", "先拿岗位编号"),
        ("吉利", "A", "座舱产品  全球适应性", "增长较强", "确认品牌与城市"),
        ("蓝思科技", "A减", "ID  项目管理  智能硬件", "转型  短期承压", "先确认事业部"),
        ("锐明技术", "B加", "商用车产品  产品市场", "收入增  利润降", "区分产品与交付"),
        ("视源股份", "B加", "交互硬件  AI教育  会议产品", "企业服务增长", "询问隐藏岗位"),
        ("快递100", "B加", "AI产品经理  小程序  Agent", "AI产品化活跃", "确认产品线"),
        ("黑白调和傲风", "B", "ID  用户体验  智能家具", "校招活跃", "验证智能化落地"),
        ("ZURU", "B减", "3D设计  新品项目", "全球化私企", "准备设计作业"),
        ("爱奇迹", "C", "硬件产品  CMF  海外研究", "信息有限", "先判断行业接受度"),
    ]
    for data in summary_rows:
        cells = table.add_row().cells
        for i, value in enumerate(data):
            cells[i].text = value
    style_table(table, widths=[1.25, 0.66, 2.04, 1.45, 1.5], header=True, font_size=8.5)

    add_heading(doc, "招聘会信息", 1)
    add_two_col_table(doc, [
        ("时间", "2026年9月19日  08时30分至17时30分"),
        ("地点", "深圳福田会展中心5号馆"),
        ("报名", "下载并注册JOBS海归App  在线简历完善度不低于60%  首页领取入场二维码"),
        ("资格", "不限毕业时间  应届生和有经验的海归均可报名"),
        ("材料", "智能汽车HMI版  智能硬件版  AI产品经理版电子简历  通用纸质简历8至10份  作品集链接与二维码"),
    ])

    for idx, c in enumerate(companies, start=1):
        doc.add_page_break()
        add_heading(doc, f"{idx}  {c['name']}", 1)
        add_two_col_table(doc, [
            ("优先级", c["priority"]),
            ("匹配判断", c["match"]),
            ("当前状态", c["status"]),
        ])
        add_heading(doc, "企业概况", 2)
        add_para(doc, c["overview"])
        add_heading(doc, "产品与服务方向", 2)
        add_bullets(doc, c["products"])
        add_heading(doc, "近期业务状态", 2)
        add_para(doc, c["current"])
        add_heading(doc, "9月19日公开岗位", 2)
        add_para(doc, c["event_jobs"])
        add_heading(doc, "岗位含义与个人匹配", 2)
        add_para(doc, c["role_read"])
        add_para(doc, c["fit"], bold_lead=None)
        add_heading(doc, "现场调研目标", 2)
        add_bullets(doc, c["goals"])
        add_heading(doc, "风险与核实重点", 2)
        add_para(doc, c["risks"])
        add_source_line(doc, c["sources"])

    doc.add_page_break()
    add_heading(doc, "其他已公开企业快速判断", 1)
    other = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["企业", "公开岗位或业务", "与当前方向的关系", "建议"]):
        other.rows[0].cells[i].text = h
    rows = [
        ("顺丰集团", "供应链管培  产品管理  行业解决方案  科技人才", "适合供应链产品或数字化产品  与HMI较远", "有明确产品经理岗位再投"),
        ("江波龙", "AI应用  IC算法  测试  FAE", "存储芯片与半导体岗位为主  专业门槛高", "不作为主投"),
        ("南方电网与深圳供电局", "高层次人才  博士后", "公开岗位层级与当前学历阶段不匹配", "仅了解后续常规校招"),
        ("驭势科技", "自动驾驶算法  集成交付  民航售前", "行业相关但缺少HMI和产品岗位", "询问隐藏产品岗位"),
        ("北芯生命", "医疗器械系统  测试  硬件  软件  创新研发", "智能硬件相关  医疗法规门槛较高", "有产品或人因岗位再投"),
        ("贝赛思与国际学校", "教师  课程协调  招生  行政", "与产品体验方向不符", "不投入主要时间"),
        ("金融与保险机构", "财富管理  理财顾问  保险营销", "常见代理制或销售导向  与目标方向不符", "除非主动转销售否则不投"),
        ("苏澜士", "产品  工业设计  AI培训  研发及高管岗位同时出现", "岗位跨度异常大且公开信息少", "先核实产品  团队  合同主体"),
    ]
    for data in rows:
        cells = other.add_row().cells
        for i, value in enumerate(data):
            cells[i].text = value
    style_table(other, widths=[1.18, 2.12, 2.05, 1.55], header=True, font_size=8.6)

    add_heading(doc, "现场统一提问清单", 1)
    questions = [
        "这个岗位的正式名称和岗位编号是什么  属于哪个事业部和产品线",
        "核心工作产出是什么  PRD  场景定义  交互方案  用户研究  项目排期  客户交付还是销售指标",
        "岗位是面向2027届校招还是社会招聘  2027年4月海外毕业是否符合",
        "工作地点和团队所在地是什么  是否需要长期出差  驻厂或海外派驻",
        "专业要求是否限制计算机  电子或机械  HCI和工业设计背景如何评估",
        "是否需要作品集  重点看研究过程  产品思考  视觉表达还是工程落地",
        "面试流程中是否有产品题  设计作业  笔试  群面或技术面",
        "现场投递与官网投递如何衔接  是否会占用投递次数",
        "预计何时开始筛选和面试  后续联系人和截止日期是什么",
    ]
    add_bullets(doc, questions)

    add_heading(doc, "展位记录模板", 1)
    record = doc.add_table(rows=10, cols=2)
    fields = ["企业与展位", "联系人", "岗位与编号", "事业部与产品线", "工作地点", "岗位核心产出", "专业与毕业时间", "作品集与面试", "投递入口与截止日期", "下一步行动"]
    for i, field in enumerate(fields):
        record.rows[i].cells[0].text = field
        record.rows[i].cells[1].text = ""
    style_table(record, widths=[1.55, 5.35], header=False, font_size=9.5)
    for row in record.rows:
        set_cell_shading(row.cells[0], PALE_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
        row.cells[1].paragraphs[0].add_run("\n")

    add_heading(doc, "建议参会顺序", 1)
    add_bullets(doc, [
        "08时30分至10时  德赛西威  荣耀  腾讯",
        "10时至12时  海尔智家  吉利  蓝思科技",
        "13时至14时30分  视源股份  锐明技术  快递100",
        "14时30分至15时30分  黑白调和傲风  ZURU  驭势科技",
        "15时30分以后  返回最感兴趣的两至三家  补问岗位编号  投递入口和面试时间",
    ])
    add_para(doc, "本报告中的经营数据以已公开定期报告和公司资料为基础。企业状态不等同于个别团队的稳定性，招聘岗位也不代表最终录用名额。现场判断应以具体岗位、事业部、工作地点、合同主体和后续流程为准。", space_after=8)

    doc.core_properties.title = "9月19日深圳海归招聘会企业调研报告"
    doc.core_properties.subject = "企业产品服务  经营状态  招聘岗位  匹配判断  现场提问"
    doc.core_properties.author = "孔维鹏"
    doc.core_properties.keywords = "招聘会, 企业调研, AI产品, 智能硬件, HMI, 用户研究"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_doc()
